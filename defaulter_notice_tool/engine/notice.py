import os
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

from . import config

# NOTE: the wording below is a generic placeholder. In the real project the
# advocate supplies the exact format, so only these two blocks need replacing.

EARLY_BODY = (
    "Our records show that you have failed to pay {missed} instalment(s) of the above "
    "loan, with an overdue amount of Rs. {overdue}. This is a final reminder to clear "
    "the dues within {days} days of receiving this notice, failing which {lender} "
    "will be compelled to take further action."
)

SEVERE_BODY = (
    "Despite repeated reminders you have defaulted on {missed} instalments of the above "
    "loan. The overdue amount is Rs. {overdue} and the principal outstanding is Rs. {pos}. "
    "Through this notice you are called upon to pay the overdue amount within {days} days "
    "of receipt, failing which my client shall initiate appropriate legal proceedings at "
    "your risk, cost and consequences."
)


def _money(value):
    try:
        return f"{float(value):,.2f}"
    except (TypeError, ValueError):
        return "N/A"


def _safe_filename(text):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", str(text)).strip("_")


def _para(doc, text="", bold=False, align=None, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def build_notice(row, stage, out_dir):
    """Write one .docx notice for one customer and return its path."""
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin = section.bottom_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    severe = stage == "severe"

    if severe:
        _para(doc, config.ADVOCATE_NAME, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=0)
        _para(doc, f"Enrolment No. {config.ADVOCATE_ENROLMENT}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        _para(doc, config.ADVOCATE_ADDRESS, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
        title = "LEGAL NOTICE"
    else:
        _para(doc, config.LENDER_NAME, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=14)
        title = "PAYMENT REMINDER NOTICE"

    _para(doc, f"Date: {date.today().strftime('%d %B %Y')}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(doc, title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, space_after=12)

    _para(doc, "To,", space_after=0)
    _para(doc, str(row.get("name", "")), bold=True, space_after=0)
    address = row.get("address")
    if isinstance(address, str) and address.strip():
        _para(doc, address.strip(), space_after=10)
    else:
        _para(doc, "", space_after=10)

    _para(doc, f"Subject: Default in repayment of loan account no. {row['loan_no']}", bold=True, space_after=10)
    _para(doc, "Sir / Madam,")

    template = SEVERE_BODY if severe else EARLY_BODY
    body = template.format(
        missed=int(row["missed"]),
        overdue=_money(row.get("overdue")),
        pos=_money(row.get("outstanding")),
        days=config.NOTICE_RESPONSE_DAYS,
        lender=config.LENDER_NAME,
    )
    p = _para(doc, body, space_after=14)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    _para(doc, "Yours faithfully,", space_after=28)
    signer = config.ADVOCATE_NAME if severe else f"Authorised Signatory, {config.LENDER_NAME}"
    _para(doc, signer, bold=True)

    os.makedirs(out_dir, exist_ok=True)
    fname = f"{_safe_filename(row['loan_no'])}_{_safe_filename(row.get('name', ''))}.docx"
    path = os.path.join(out_dir, fname)
    doc.save(path)
    return path
